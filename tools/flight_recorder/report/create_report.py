from pathlib import Path
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import json, hashlib, zipfile
R=Path(os.environ.get("FLIGHT_REPORT_DIR", ".")).resolve()
d=Document()
sec=d.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.3)
for n in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3","Caption"]:
 s=d.styles[n]; s.font.name="Noto Sans CJK SC"; s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),"Noto Sans CJK SC")
 s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
normal=d.styles["Normal"]; normal.font.size=Pt(11); normal.paragraph_format.line_spacing=1.2
for n,size in [("Title",25),("Heading 1",18),("Heading 2",13),("Heading 3",11.5)]:
 d.styles[n].font.size=Pt(size);d.styles[n].font.bold=True
d.styles["Title"].font.color.rgb=RGBColor(0,0,0)
d.styles["Caption"].font.size=Pt(9); d.styles["Caption"].font.italic=False; d.styles["Caption"].font.bold=False
h=sec.header.paragraphs[0]; h.text="推理 Logits 记录与离线分析"; h.style=d.styles["Caption"]
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run("设计说明书  ·  ")
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
def p(t,style=None):
 x=d.add_paragraph(t,style);return x
def heading(t,l=1):return d.add_heading(t,l)
def page(): d.add_page_break()
def fig(n,caption,w=7.05):
 x=d.add_paragraph();x.paragraph_format.keep_with_next=True
 x.add_run().add_picture(str(R/'assets'/n),width=Inches(w))
 p(caption,"Caption")
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 trPr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trPr.append(rep)
 for row in rows:
  for c,val in zip(t.add_row().cells,row):c.text=str(val)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();nr=OxmlElement('w:cantSplit');pr.append(nr)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge,v in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');m.append(el)
   tcPr.append(m)
   if ri==0 or ri%2==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8F0F4' if ri==0 else 'F5F7F9');tcPr.append(sh)
   for pa in c.paragraphs:
    pa.paragraph_format.space_after=Pt(2);pa.paragraph_format.space_before=Pt(1);pa.paragraph_format.line_spacing=1.08
    for run in pa.runs:run.font.size=Pt(9);run.font.bold=ri==0
 p("")
 return t
def code(t):
 x=p(t);x.paragraph_format.line_spacing=1.0;x.paragraph_format.space_after=Pt(8)
 for r in x.runs:r.font.name="DejaVu Sans Mono";r.font.size=Pt(8)
def note(t):p(t,"Caption")

p("推理 Logits 记录\n与离线分析设计说明书","Title")
p("Qwen3.6-27B W8A8 · vLLM-Ascend 0.23.0\nTP=2 · MTP=3 · 图模式","Subtitle")
p("验证日期  2026-09-07\n执行环境  41_7 / vllm-023-logits-save")
heading("1. 背景介绍")
p("中文推理中偶发的语言切换、提前结束和内容错误难以稳定复现。仅保存最终文本，无法判断第一个异常 token 出现时，模型倾向是什么、采样处理改变了什么，以及系统实际选择了哪个候选。")
p("本方案覆盖测试中全部请求的全部生成位置，保存两阶段 Top-K logits、实际输出 token、EOS 信息和请求上下文。发现异常后，可以按请求与生成位置回溯当时的目标模型分布、采样结果、MTP 来源及停止原因。")
heading("1.1 目标与诊断边界",2)
p("记录范围包含 prefill 后第一个输出、MTP 接受 / 替换 / bonus token 和 EOS。输入 token IDs 完整保存一次，不额外计算输入每个位置的 prompt logits；停止后的推测尾部单独标识。")
p("低概率或非 Top-1 采样只是排查线索，不能直接判定内容错误；高概率也不能证明语义正确。Top-K 不包含完整长尾，也不足以仅凭 seed 精确重放随机采样与 MTP 验证过程。")
heading("1.2 已验证的结果",2)
table(["项目","结果"],[
["记录与对齐","95 个请求，22,004 个最终输出 token；22,392 个采样位置全部审计通过"],
["端到端开销","Top128：TPOT 中位数增加 0.84%～3.15%；Top1024：增加 4.47%～5.54%"],
["存储量级","每百万采样位置约 2.75 GB / 7.36 GB；另计共用 tokenizer"],
["验证范围","指定模型、TP2、MTP3、C1/C4、固定公开数据子集；不外推其他负载"]
],[1.25,5.8])
note("文档内容依据已验证的实现、运行清单和原始测试记录整理。配图为 Technical Whitepaper Infographic；HTML 图片为真实日志页面截图。")

page()
heading("2. 整体方案")
heading("2.1 推理与后台记录分工",2)
p("采用 cpu_pool 后端。主推理路径向独立 NPU 槽保存 logits 快照与元信息；复制线程和独立写入进程负责传输、统计及持久化。磁盘仅保存 Top-K 摘要与必要诊断字段。")
fig("fig1.png","图 1  推理主路径与后台记录路径。raw / final 分别对应约束前与目标模型采样处理后。")
p("raw 的采集点位于 grammar mask 之前；final 在目标模型采样处理之后采集。记录器保留原有温度、top-k / top-p、约束及采样逻辑，不通过记录过程改变生成规则。")
p("后台读取完整捕获域，计算 Top-K、完整分布 logsumexp、所选 token 和 EOS 的值与排名。这样既减少磁盘存储，又能避免只在 Top-K 内归一化所造成的概率高估。")
p("TP rank 0 负责记录，避免两个 rank 重复落盘。完整快照只在有界槽池中短期驻留；池满后产生背压，不静默丢弃、不自动降低 K。")
heading("容量与性能的取舍",2)
p("池化减少反复分配，并将统计和写盘移出主路径，但快照、D2H 传输、共享区复制及排队仍有成本。容量规划应同时检查 NPU 内存、主机内存、后台吞吐和磁盘保留周期。")

page()
heading("2.2 有界槽池与持久化确认",2)
fig("fig2.png","图 2  槽池生命周期：持久化确认之前，任何生产者都不能覆盖该槽。")
p("记录池预分配 8 个槽，每槽容纳本批最多 16 行 logits。快照使用独立存储，采样器后续对原始 logits 的原地修改不会破坏已捕获的 raw。")
p("复制线程等待 NPU 事件，将有效行复制到可复用的固定页 CPU 缓冲，再放入 mmap 共享槽。写入进程配置 4 个 CPU 统计线程；raw / final 数值相同时复用统计结果。")
p("二进制块完成 fsync，且 SQLite 索引事务提交之后，写入端才发送 ACK。生产端在 ACK 回调中归还槽，形成“写入快照 → 后台处理 → 持久化 → 槽复用”的闭环。")
heading("背压与失败处理",2)
p("后台处理不足时，生产者等待空闲槽；等待会反映到推理延迟中。复制或写入失败会显式报告并留下失败标记。HTTP 返回成功不能代替日志完整性审计。")
p("正常退出需要排空记录队列。异常强杀仍可能损失尚未持久化的尾部，因此离线工具必须检查关闭标记、数据校验和输出位置连续性，并将不完整日志标注出来。")

page()
heading("2.3 MTP 与最终输出对齐",2)
fig("fig3.png","图 3  MTP=3 时的两类输出示例。记录按实际输出槽映射到相应 target 行。")
p("一次 MTP 验证可能输出多个 token。全部接受时，三个 draft token 后可附加 bonus；部分接受时，记录已接受前缀和 replacement，不能把被拒绝或丢弃的 draft 槽误当成最终输出。")
p("日志保存草稿 IDs、数量、实际输出槽、普通 / accepted_draft / replacement / bonus 来源。输出处理器另存最终提交序列及停止事件，两条证据链用于逐位置对齐。")
p("因停止条件而未提交的推测尾部可以保留，但默认不参与 HTML 的最终输出筛选。审计区分“发生过采样的位置”和“实际提交给用户的位置”。")
heading("目标分布的解释范围",2)
p("MTP 页面中的概率来自 target 模型。完整 draft 概率、验证 RNG 和 replacement 的条件抽样分布没有全部保存，因此这些日志支持定位与解释实际输出，不承诺精确重放接受 / 拒绝随机过程。")
p("greedy 场景中，归一化概率描述模型偏好；实际选择仍是确定性 argmax。并发随机请求即使使用相同 seed，也不能据此保证逐 token 完全重现。")

page()
heading("2.4 内存与磁盘容量估算",2)
fig("fig4.png","图 4  固定槽池容量与长期日志容量分开规划。GB 为十进制，MiB 为二进制。",6.8)
p("单类缓冲 = 8 槽 × 16 行 × 248,320 词表 × 2 阶段 × 4 字节 = 254,279,680 字节，即 242.5 MiB。三类主要载荷合计 727.5 MiB，其中 rank 0 额外 NPU 快照为 242.5 MiB；这是分配形状推算，不是 RSS 或峰值实测。")
table(["采样位置数","Top128 日志","Top1024 日志"],[
["10 万","约 0.275 GB","约 0.736 GB"],["100 万","约 2.75 GB","约 7.36 GB"],["1,000 万","约 27.5 GB","约 73.6 GB"]
],[2.35,2.35,2.35])
p("估算采用本次样本平均 2.75 KB / 7.36 KB 每位置，已含索引、事件及压缩块开销。位置数可按“请求数 × 平均输出 token 数 ×（1 + 推测尾部余量）”估计。")
p("例如每日 100 万采样位置、保留 7 天并预留 30% 容量：Top128 约需 25.0 GB，Top1024 约需 67.0 GB；另计共用 tokenizer 22.9 MB 及超出本次输入长度的元数据增量。长提示词需单独预算，不能把样本压缩率视为保证。")

page()
heading("3. 详细设计")
heading("3.1 软件环境与运行配置",2)
table(["项","已验证配置"],[
["模型","Qwen3.6-27B W8A8；/mnt/weights/Qwen3.6-27B-w8a8"],
["镜像","quay.io/ascend/vllm-ascend:v0.23.0-openeuler"],
["框架","vLLM 0.23.0+empty；vLLM-Ascend 0.23.0"],
["运行库","Python 3.12.13；PyTorch 2.10.0+cpu；torch-npu 2.10.0.post4"],
["设备软件","CANN 9.1.0；驱动报告 26.0.rc1"],
["并行 / 图","TP=2，MTP=3；target FULL_DECODE_ONLY，draft FULL"],
["容量 / 调度","max_model_len=4096；max_num_seqs=4；内存利用率 0.75；关闭 prefix caching"],
["运行边界","41_7 / vllm-023-logits-save；仅使用本任务容器与授权持久化目录"]
],[1.15,5.9])
heading("3.2 主要修改文件",2)
note("下表路径分别相对于 vLLM 与 vLLM-Ascend worktree 根目录。")
table(["仓库 / 文件","主要代码改动"],[
["vLLM · vllm/v1/sample/flight_recorder.py","有界 IPC、独立写入进程、共享快照统计、索引 / 校验 / ACK、完整采样参数事件"],
["vLLM · vllm/v1/sample/sampler.py","采样处理后回调，保留 greedy 分支语义"],
["vLLM · vllm/v1/sample/ops/topk_topp_sampler.py","native 随机路径截断后回调"],
["vLLM · vllm/v1/engine/output_processor.py","输入、最终提交 token、停止及取消事件"],
["Ascend · vllm_ascend/sample/flight_recorder.py","采集入口与后端选择"],
["Ascend · vllm_ascend/sample/flight_recorder_pool.py","NPU / 固定页 CPU / 共享槽池、事件、复制线程及槽归还"]
],[3.5,3.55])

page()
heading("3.2 主要修改文件（续）",2)
table(["Ascend 文件","主要代码改动"],[
["vllm_ascend/sample/sampler.py","随机采样处理后的快照点"],
["vllm_ascend/sample/rejection_sampler.py","bonus / target 行映射及 target 处理后快照"],
["vllm_ascend/worker/model_runner_v1.py","grammar 前开始、采样后提交、实际图模式证据"],
["vllm_ascend/worker/worker.py","退出时排空并关闭记录池"],
["tools/flight_recorder/analyze.py","列表 / 查询 / 审计、tokenizer 校验和 HTML 导出"],
["tools/flight_recorder/viewer.html","离线界面、概率筛选、位置导航、候选表"],
["tools/flight_recorder/benchmark.py","固定公开用例、自然 EOS、并发预热、流式计时"],
["tools/flight_recorder/manage.py","任务进程组管理、退出检查、临时槽清理"],
["tools/flight_recorder/test_recorder.py","数值 / MTP / 异步复用 / 失败 / 生命周期回归"],
["tools/flight_recorder/profile_components.py","组件耗时穿刺"],
["tools/flight_recorder/run_pool_matrix.py；summarize.py","测试编排和结果汇总"]
],[3.65,3.4])
heading("3.3 关键代码与同步边界",2)
p("model_runner_v1.py 的 begin 位于 grammar 修改前，finish 位于原有 _sample 返回后。以下节选展示关键接口及数据所有权：")
code("flight_capture = begin(self, logits, spec_decode_metadata)\n# 原有 grammar 与 _sample 流程\nif flight_capture is not None:\n    flight_capture.finish(sampler_output)\n\nslot.raw[:self.nrows].copy_(logits)\nself.slot.final.index_copy_(0, rows, logits.to(torch.float32))")
p("复制线程单独使用 @torch.inference_mode()，因为该状态为线程局部。提交写入任务时注册 on_done 回调，只有持久化 ACK 到达后才 release(slot)，防止异步消费者读取到下一批覆盖的数据。")

page()
heading("3.4 日志组织与字段",2)
code("run/manifest.json\nrun/trace/PID.sqlite\nrun/trace/PID.bin\nrun/trace/PID.closed\nrun/trace/FAILED-*\ntokenizer/                  # 多个 run 共用的 tokenizer 快照")
p("每个生产进程对应一个 shard。SQLite 保存 request_id、position、token_id、chunk 和 row；PID.bin 顺序追加 NPZ / DEFLATE 无损压缩块。块索引包含偏移、字节数及 SHA256。")
table(["记录层级","必要字段"],[
["生成位置","raw / final 的 Top-K ID（int32）和 logits（FP32）；实际输出与 EOS 的值和全词表排名"],
["分布统计","完整捕获域 logsumexp；finite、NaN、+Inf、−Inf 数量；final_seen"],
["MTP / 批次","草稿 IDs / 数量、输出槽和来源；请求映射、时间、实际图模式及策略标记"],
["请求事件","实际输入 token IDs、生效 SamplingParams 全部字段（含默认值）"],
["输出事件","最终提交 token 序列、finish_reason、stop_reason"],
["来源与解码","manifest 中的参数、环境及源码指纹；tokenizer 文件哈希"]
],[1.2,5.85])
heading("3.5 概率与排名定义",2)
p("概率按完整捕获域归一化：p(token) = exp(logit(token) − logsumexp)。Top-K 覆盖率为所保留候选概率之和；概率差为同一步 p(Top-1) − p(Top-2)。不得只对保存的 Top-K 做 softmax。")
p("CPU 使用 FP64 稳定计算 logsumexp。该数值统计不承诺逐位复现设备 softmax 舍入。实际输出及配置 EOS 即使不在 Top-K 中，也单独保留其值与排名。")
p("采样排名采用“严格大于该 logit 的候选数 + 1”。并列最高概率候选的排名均为 1；候选表中的展示序号与这种排名含义不同。未知概率不会匹配阈值筛选。")
p("导出前校验 tokenizer.json、tokenizer_config.json、vocab.json，防止用另一套 tokenizer 解码历史 token ID。Top-K 截断的长尾无法在事后恢复，应结合覆盖率决定是否提高 K。")

page()
heading("3.6 HTML 离线分析界面",2)
p("三类条件可选 raw / final，并按 AND 组合：采样排名 > 1、所选 token 概率小于阈值、同一步 Top-1 与 Top-2 概率差小于阈值。默认只看最终输出，箭头在命中位置间导航。")
fig("ui-filters.png","图 5a  真实页面筛选：final、排名 > 1、概率 < 0.5、Top1−Top2 差 < 0.1，命中 1 / 150 个位置。",6.8)
table(["筛选条件","本例设置","含义"],[
["分布阶段","final","按目标模型采样处理后的分布计算"],
["采样排名","大于 1","所选候选不是并列最高概率"],
["所选概率","小于 0.5","所选 token 在完整捕获域中的概率"],
["候选差值","小于 0.1","同一步 Top-1 与 Top-2 的概率差"]
],[1.6,1.3,4.15])
p("条件同时成立时，页面只保留命中位置；本例为 1 / 150。可使用箭头跳转，也可输入任意生成位置直接查看。raw / final 切换有助于分清模型原始偏好与采样处理的影响。")
page()
heading("3.6 单位置候选分布分析",2)
fig("ui-detail.png","图 5b  同一请求位置 75：“蛋”，accepted_draft，排名 2；final 概率 0.455475，Top1−Top2 差 0.089049。",6.8)
table(["位置 75 的统计","raw","final"],[
["实际输出 token","蛋（ID 97457）","蛋（ID 97457）"],
["来源与排名","accepted_draft；排名 2","accepted_draft；排名 2"],
["所选 token 概率","0.465522","0.455475"],
["Top-1 与 Top-2 概率差","0.061984","0.089049"],
["Top-K 概率覆盖率","0.999825","1.000000"]
],[2.55,2.25,2.25])
p("页面将实际输出候选以浅黄色标出，同时给出 EOS 候选的值与排名。此例所选 token 并非最高概率，但这只是一次正常随机推理的可回溯证据，不能据此判定输出有误。")
note("截图来自公开 MGSM 用例的实际日志，证明查询可用；不代表复现了用户原先的偶发故障。HTML 无 CDN 依赖，导出后可直接离线打开。")

page()
heading("3.7 开启、查询与退出",2)
p("记录器 TOPK 控制保存规模，与生成 SamplingParams 中的 top_k 相互独立。未设置记录目录时关闭记录。每次运行应使用独立目录：")
code("export VLLM_FLIGHT_RECORDER_DIR=/path/to/new-run/trace\nexport VLLM_FLIGHT_RECORDER_BACKEND=cpu_pool\nexport VLLM_FLIGHT_RECORDER_TOPK=128\nexport VLLM_FLIGHT_RECORDER_PROBES=248044,248046\nexport VLLM_FLIGHT_RECORDER_QUEUE=8")
p("启动服务显式设置 --additional-config '{\"enable_cpu_binding\":false}' 和 --shutdown-timeout 30。manage.py 检查整个已验证任务进程组退出，并仅清理对应已退出创建者的临时共享槽。")
p("enable_reduce_sample 明确不支持，避免把局部词表误认为完整词表。PP / DP、多模态和可续接流式输入未验证，不能扩大兼容性结论。")
heading("离线排查步骤",2)
p("先运行 audit，确认数据校验、索引 / 载荷一致、生成位置连续、final_seen 及最终输出逐 token 对齐；再定位第一个异常位置，对比 raw / final、EOS、输出来源和停止原因。")
code("python analyze.py /path/to/run/trace list\npython analyze.py /path/to/run/trace audit\npython analyze.py /path/to/run/trace query \\\n  --request REQUEST_ID --position 75\npython analyze.py /path/to/run/trace html \\\n  --request REQUEST_ID --tokenizer /path/to/tokenizer \\\n  --out result.html")
p("查询仅需 Python 与 NumPy；HTML 导出还需匹配的本地 tokenizer 和 Transformers。按单请求导出，避免把整个压缩日志集展开为大量文本 JSON。")
heading("工程目录与版本",2)
note("容器内 worktree 根：/home/w00498770/dev/worktrees/vllm-023-logits-save/")
code("vllm/\n  a20a56bbb60d511876d8879b60de1752e2c5a8fc\nvllm-ascend/\n  389403d0ba0a6f9945ea166fd91c8156d7725da2\nbranch: feat/logits-flight-recorder")
note("持久化结果：/home/w00498770/dev/artifacts/vllm-023-logits-save/。各 run 的 manifest 是实际测试参数及来源快照依据；源码交付版本包含工具、退出管理及文档。")

page()
heading("5. 穿刺结果")
heading("5.1 数据与测量口径",2)
p("使用 MGSM 中文前 8 题与 GSM8K 英文前 8 题，分别在 C1 / C4 运行。MGSM 是 GSM8K 的翻译子集，不能视为独立题库；此测试属于固定子集机制穿刺，不是官方全量评测。")
p("贪心配置为 temperature=0、seed=42、输出上限 512、enable_thinking=false，使用自然 EOS，并按对应并发预热。随机验证使用两个数据集各 2 题，temperature=0.7、top_p=0.9、top_k=20，覆盖 C1 / C4。")
p("TPOT 为客户端首末输出 token 块间隔除以输出间隔数的有效平均值，表中再取请求中位数。MTP 可一次返回多个 token，因此该指标不等于内部逐 token ITL。未取得独占 CPU 租约，结果只代表本次环境。")
heading("5.2 关闭记录与启用记录的对照",2)
table(["记录配置","并发","TPOT 中位数\nms","相对关闭\n记录","输出 token"],[
["关闭记录","1","9.602","—","4,039"],
["cpu_pool Top128","1","9.905","+3.15%","4,039"],
["cpu_pool Top1024","1","10.032","+4.47%","4,039"],
["关闭记录","4","14.474","—","4,115"],
["cpu_pool Top128","4","14.595","+0.84%","4,115"],
["cpu_pool Top1024","4","15.277","+5.54%","4,115"]
],[2.0,.55,1.6,1.45,1.45])
p("各 Top-K / 并发贪心组合均有 16 / 16 请求与对应关闭记录对照逐 token 相同。MTP 接受 / 提议计数也一致：C1 为 2,873 / 3,525，C4 为 2,927 / 3,579。")
p("每组数值答案匹配 15 / 16，另 1 条达到长度上限；length 停止不能算自然 EOS。总吞吐还受首批延迟和调度波动影响，不能将某次吞吐较高解释为记录器加速了模型。")
heading("5.3 主路径与背压观察",2)
p("主线程记录钩子 CPU 耗时中位数约 0.59～0.60 ms / 批；后台 CPU 统计约 10 ms / 批。取槽等待中位数约 0.008～0.009 ms，Top1024 最大观察值为 52.44 ms / 批。池化降低常态开销，但仍存在背压尾延迟。")

page()
heading("5.4 完整性、存储与覆盖率",2)
table(["配置","请求","采样位置","日志 MB","字节 / 位置"],[
["Top128","41","10,150","27.95","2,754"],
["Top1024","54","12,242","90.16","7,365"]
],[1.4,.7,1.55,1.55,1.85])
p("合计 95 个请求（含预热）、22,004 个最终提交 token、22,392 个采样位置，全部审计通过；差额为停止后的推测尾部。普通、accepted_draft、replacement、bonus 来源及 NONE / FULL 图模式均在记录中实际出现。")
p("Top128 的 raw 覆盖率最小 98.5742%、5 分位 99.9866%；Top1024 最小 99.5382%、5 分位 99.9959%。覆盖率统计含预热和尾部，不是未来请求的下界保证；两组样本构成不同，也不能将磁盘比值视为只由 K 决定。")
heading("5.5 正确性与组件穿刺",2)
p("15 项最终回归测试通过，覆盖 CPU 持久化 / 损坏 / 失败、NPU 全词表对照、MTP 变长映射、即时覆盖与背压、跨线程 inference_mode、完整默认参数及任务进程组清理。")
p("4 × 248,320 FP32 的独立组件测量中，NPU 快照复制中位数 0.046 ms，单份完整 D2H 至固定页 CPU 约 0.285 ms，4 线程 CPU 摘要约 10.1 ms。组件值不能相加替代端到端 TPOT。")
p("浏览器验证覆盖三类筛选、AND 组合、raw / final 切换、空结果、清除及位置导航。示例中的非 Top-1 位置有 3 个；概率 < 0.2 可定位位置 68；图 5 组合条件定位位置 75。")
heading("5.6 复核命令与证据",2)
code("# 在目标容器的 vllm-ascend/tools/flight_recorder 中\nRUN_NPU_TESTS=1 ASCEND_RT_VISIBLE_DEVICES=2 \\\n  python -m pytest -q test_recorder.py\npython analyze.py /path/to/pooled128/trace audit\npython analyze.py /path/to/pooled1024/trace audit\npython summarize.py")
note("原始证据：results.json、各 run 的 manifest / metrics / audit、unit-tests-delivery.log、npu-final-release.txt。推理进程与本任务 NPU 已核对释放；离线文档制作不启动模型。")
heading("数据与源码来源",2)
note("MGSM：https://github.com/google-research/url-nlp/tree/main/mgsm\nGSM8K：https://github.com/openai/grade-school-math/tree/master/grade_school_math/data\n源码：https://github.com/wanghuanjun2113/vllm\n源码及设计文档：https://github.com/wanghuanjun2113/vllm-ascend\n两仓库分支均为 feat/logits-flight-recorder；精确提交见第 3.7 节。")
d.core_properties.title="推理 Logits 记录与离线分析设计说明书"
d.core_properties.subject="Qwen3.6-27B W8A8 / vLLM-Ascend / TP2 / MTP3"
d.core_properties.author="工程设计"
d.core_properties.keywords="Logits,Top-K,MTP,TPOT,离线分析"
for st in d.styles:
 for el in list(st._element.iter(qn("w:pBdr"))):
  el.getparent().remove(el)
for pa in d.paragraphs:
 for el in list(pa._p.iter(qn("w:pBdr"))):
  el.getparent().remove(el)
out=R/"logits-design.docx";d.save(out)
assert zipfile.is_zipfile(out)
print(json.dumps({"path":str(out),"bytes":out.stat().st_size,"sha256":hashlib.sha256(out.read_bytes()).hexdigest(),"images":len(d.inline_shapes)},ensure_ascii=False))
