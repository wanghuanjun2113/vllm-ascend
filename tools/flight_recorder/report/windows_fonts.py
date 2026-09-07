from pathlib import Path
import os
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import json,hashlib,zipfile
R=Path(os.environ.get("FLIGHT_REPORT_DIR", ".")).resolve()
d=Document(R/"logits-design-formatted.docx")
def setfonts(rpr,code=False):
 f=rpr.find(qn("w:rFonts"))
 if f is None:f=OxmlElement("w:rFonts");rpr.insert(0,f)
 f.attrib.clear()
 for k,v in {"ascii":"Consolas" if code else "Arial","hAnsi":"Consolas" if code else "Arial","eastAsia":"Microsoft YaHei","cs":"Arial","hint":"default"}.items():f.set(qn("w:"+k),v)
 lang=rpr.find(qn("w:lang"))
 if lang is None:lang=OxmlElement("w:lang");rpr.append(lang)
 lang.set(qn("w:val"),"en-US");lang.set(qn("w:eastAsia"),"zh-CN")
for st in d.styles:
 if st.type==1 or st.type==2:
  setfonts(st._element.get_or_add_rPr())
for e in d.styles.element.findall(".//"+qn("w:rPrDefault")+"/"+qn("w:rPr")):setfonts(e)
paragraphs=list(d.paragraphs)
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=Pt(14)
    for r in p.runs:r.font.size=Pt(9.5)
    paragraphs.append(p)
for s in d.sections:
 paragraphs.extend(s.header.paragraphs);paragraphs.extend(s.footer.paragraphs)
for p in paragraphs:
 pp=p._p.get_or_add_pPr()
 op=pp.find(qn("w:overflowPunct"))
 if op is None:op=OxmlElement("w:overflowPunct");pp.append(op)
 op.set(qn("w:val"),"0")
 for r in p.runs:
  old=r.font.name or ""
  iscode=("Mono" in old or old=="Consolas")
  setfonts(r._r.get_or_add_rPr(),iscode)
  if iscode:r.font.size=Pt(8.5);p.paragraph_format.line_spacing=Pt(12)
 if p.style.name=="Subtitle":
  for r in p.runs:r.italic=False;r.font.size=Pt(12)
d.styles["Normal"].paragraph_format.line_spacing=Pt(18)
for p in d.paragraphs:
 if p._p.xpath(".//w:drawing"):
  p.paragraph_format.line_spacing=1.0
 else:
  p.paragraph_format.right_indent=Pt(11)
# Explicit typography in the theme prevents fallback to Mac-only defaults.
for part in d.part.package.parts:
 if str(part.partname).startswith("/word/theme/"):
  root=etree.fromstring(part.blob);ns={"a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
  for group in root.xpath("//a:majorFont|//a:minorFont",namespaces=ns):
   for tag,font in [("latin","Arial"),("ea","Microsoft YaHei"),("cs","Arial")]:
    e=group.find("{"+ns["a"]+"}"+tag)
    if e is not None:e.set("typeface",font)
   for e in group.findall("{"+ns["a"]+"}font"):
    if e.get("script")=="Hans":e.set("typeface","Microsoft YaHei")
  part._blob=etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
# Record the requested families in the Word font table.
for part in d.part.package.parts:
 if str(part.partname)=="/word/fontTable.xml":
  root=etree.fromstring(part.blob)
  for e in list(root):root.remove(e)
  for family,pitch,kind in [("Arial","variable","swiss"),("Microsoft YaHei","variable","swiss"),("Consolas","fixed","modern")]:
   e=OxmlElement("w:font");e.set(qn("w:name"),family)
   f=OxmlElement("w:family");f.set(qn("w:val"),kind);e.append(f)
   f=OxmlElement("w:pitch");f.set(qn("w:val"),pitch);e.append(f)
   root.append(e)
  part._blob=etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
for f in list(d.styles.element.iter(qn("w:rFonts"))):
 setfonts(f.getparent(), f.get(qn("w:ascii"))=="Consolas")
out=R/"logits-design-windows.docx";d.save(out)
with zipfile.ZipFile(out) as z:
 for name in ["word/document.xml","word/styles.xml"]:
  root=etree.fromstring(z.read(name))
  for f in root.iter(qn("w:rFonts")):
   assert not any("Theme" in k for k in f.attrib)
   assert f.get(qn("w:eastAsia"))=="Microsoft YaHei"
   assert f.get(qn("w:ascii")) in ["Arial","Consolas"]
 old=zipfile.ZipFile(R/"logits-design-formatted.docx")
 for n in old.namelist():
  if n.startswith("word/media/"):assert old.read(n)==z.read(n)
check=Document(out)
for t in check.tables:
 widths=[c.width.twips for c in t.columns];assert sum(widths)==10080
 for row in t.rows:assert [c.width.twips for c in row.cells]==widths
meta={"path":str(out),"bytes":out.stat().st_size,"sha256":hashlib.sha256(out.read_bytes()).hexdigest(),"fonts":{"chinese":"Microsoft YaHei","latin":"Arial","code":"Consolas"},"tables":len(check.tables),"images":len(check.inline_shapes),"native_windows_rendered":False}
(R/"windows-font-check.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2))
print(json.dumps(meta,ensure_ascii=False))
